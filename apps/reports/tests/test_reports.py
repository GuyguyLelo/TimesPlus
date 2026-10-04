from io import BytesIO
from datetime import date, time

from django.contrib.admin.sites import site
from django.test import RequestFactory, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.audit.admin import AuditLogAdmin
from apps.audit.models import AuditLog
from apps.overtime.models import OvertimeRequest
from apps.overtime.services.declaration import enregistrer_declaration
from apps.overtime.tests.helpers import PASSWORD, build_referential, make_user
from apps.reports.models import GeneratedReport
class ReportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.world = build_referential()
        cls.admin = make_user("admin", "SUPER_ADMIN", superuser=True)

    def _approve(self, day, start, end):
        demande, _result = enregistrer_declaration(
            agent=self.world["agent"],
            date_travail=day,
            heure_debut=start,
            heure_fin=end,
            motif="Travail supplémentaire",
            user=self.world["agent_user"],
        )
        return demande

    def test_excel_pdf_and_totals(self):
        self._approve(date(2026, 10, 7), time(17, 0), time(20, 0))
        self._approve(date(2026, 10, 8), time(17, 0), time(20, 30))
        self.client.login(username="admin", password=PASSWORD)
        query = "?debut=2026-10-01&fin=2026-10-31&statut=APPROUVE"
        excel = self.client.get(reverse("reports:excel") + query)
        self.assertEqual(excel.status_code, 200)
        self.assertIn("spreadsheetml", excel["Content-Type"])
        workbook = load_workbook(BytesIO(excel.content))
        self.assertEqual(workbook.sheetnames, ["Synthèse", "Détails", "Par service"])
        summary = workbook["Synthèse"]
        labels = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(1, 15)}
        self.assertEqual(labels["Total minutes"], 390)
        self.assertEqual(labels["Nombre de demandes"], 2)

        pdf = self.client.get(reverse("reports:pdf") + query + "&kind=ADMINISTRATIF")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertTrue(GeneratedReport.objects.filter(format=GeneratedReport.Format.PDF).exists())

    def test_service_filter_limits_the_total(self):
        self._approve(date(2026, 10, 7), time(17, 0), time(20, 0))
        other, _result = enregistrer_declaration(
            agent=self.world["other"],
            date_travail=date(2026, 10, 8),
            heure_debut=time(17, 0),
            heure_fin=time(18, 0),
            motif="Autre service",
            user=self.world["rh_user"],
        )
        other.statut = OvertimeRequest.Statut.APPROUVE
        other.save(update_fields=["statut"])
        self.client.login(username="rh", password=PASSWORD)
        response = self.client.get(
            reverse("reports:excel") + f"?debut=2026-10-01&fin=2026-10-31&service={self.world['dsi'].pk}"
        )
        workbook = load_workbook(BytesIO(response.content))
        summary = workbook["Synthèse"]
        labels = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(1, 15)}
        self.assertEqual(labels["Total minutes"], 180)

    def test_audit_cannot_be_deleted_from_admin(self):
        from django.contrib.admin.sites import site

        model_admin = site._registry[AuditLog]
        self.assertIsInstance(model_admin, AuditLogAdmin)
        request = RequestFactory().get("/admin/audit/auditlog/")
        request.user = self.admin
        self.assertFalse(model_admin.has_delete_permission(request))
